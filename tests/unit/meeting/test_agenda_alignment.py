"""tests/unit/meeting/test_agenda_alignment.py -- 议程与转写实录时间线对齐引擎单元测试。"""
from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock, patch
import docx
import pytest

from domains.meeting.tasks.agenda_minutes.agenda_parser import (
    AgendaItemParsed,
    AgendaPlan,
    clean_presenter_names,
    parse_agenda_text,
)
from domains.meeting.tasks.agenda_minutes.alignment_engine import (
    align_agenda_with_transcript,
    group_transcript_blocks,
)
from domains.meeting.tasks.agenda_minutes.steps.agenda_minutes_agent import (
    AgendaMinutesAgent,
)
from domains.meeting.tasks.agenda_minutes.steps.agenda_minutes_render import (
    AgendaMinutesRender,
)
from infra.exporters.html.agenda_minutes import (
    format_agenda_minutes_markdown,
    render_agenda_minutes_html,
)

PROJECT_ROOT = Path(__file__).resolve().parent
while PROJECT_ROOT.parent != PROJECT_ROOT and not (PROJECT_ROOT / "domains").exists():
    PROJECT_ROOT = PROJECT_ROOT.parent

AGENDA_DIR = (
    PROJECT_ROOT / "data" / "1" / "agenda"
    if (PROJECT_ROOT / "data" / "1" / "agenda").exists()
    else PROJECT_ROOT / "agenda"
)

def test_alignment_unmentioned_agenda_marks_skipped():
    """核心防幻觉单测：验证 Test 1 中发言块独立路由。
    
    结果：
    - 议题01（小艺慧记，申家坤）：申家坤未参会未发言，如实判为 skipped，块数为 0；
    - 议题02（翻译海外，刘畅）：正常讨论，为 discussed；
    - 议题03（HAG）：未参会，为 skipped；
    - 议题04（SpeechASR，陆敬怡等）：陆敬怡发言准确归入议题04（不再被孤立标题抢占到议题01），为 discussed。
    """
    docx_path = AGENDA_DIR / "test1" / "商评1.docx"
    txt_path = AGENDA_DIR / "test1" / "商评1.txt"
    if not docx_path.is_file() or not txt_path.is_file():
        pytest.skip("Test 1 files not found")

    doc = docx.Document(str(docx_path))
    transcript = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    plan = parse_agenda_text(txt_path.read_text(encoding="utf-8"))

    res = align_agenda_with_transcript(plan, transcript)
    assert len(res.alignments) == 4
    # 核心判决：02 与 04 有效讨论，01 与 03 确实未讨论
    assert res.discussed_count == 2
    assert res.skipped_count == 2

    a1, a2, a3, a4 = res.alignments
    assert a1.item.seq == "01"
    assert a1.status == "skipped"
    assert a1.matched_blocks == []

    assert a2.item.seq == "02"
    assert a2.status == "discussed"
    assert "刘畅" in a2.matched_speakers
    assert len(a2.matched_blocks) > 30

    assert a3.item.seq == "03"
    assert a3.status == "skipped"
    assert a3.matched_blocks == []

    assert a4.item.seq == "04"
    assert a4.status == "discussed"
    assert "陆敬怡" in a4.matched_speakers
    assert len(a4.matched_blocks) > 100

    # 验证按现场研讨真实流向排序：先讨论的 02，随后 04，最后是 skipped 的 01 和 03
    chrono = res.chronological_alignments
    assert len(chrono) == 4
    assert [c.item.seq for c in chrono] == ["02", "04", "01", "03"]

def test_alignment_presenter_name_grounding():
    """验证 Test 2 中 9 项议题依靠汇报人主权发言彻底解决标题错位问题。"""
    docx_path = AGENDA_DIR / "test2" / "商评2.docx"
    txt_path = AGENDA_DIR / "test2" / "商评2.txt"
    if not docx_path.is_file() or not txt_path.is_file():
        pytest.skip("Test 2 files not found")

    doc = docx.Document(str(docx_path))
    transcript = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    plan = parse_agenda_text(txt_path.read_text(encoding="utf-8"))

    res = align_agenda_with_transcript(plan, transcript)
    assert len(res.alignments) == 9
    assert res.discussed_count == 9

    align_map = {a.item.seq: a for a in res.alignments}

    # 郑爽准确回到 04 IDS，而不是留在小艺慧记或 HAG
    assert "郑爽" in align_map["04"].matched_speakers
    assert len(align_map["04"].matched_blocks) > 100

    # 陈啟锴在 05 HAG
    assert "陈啟锴" in align_map["05"].matched_speakers
    assert len(align_map["05"].matched_blocks) > 100

    # 陆敬怡在 06 SpeechASR
    assert "陆敬怡" in align_map["06"].matched_speakers

    # 孙鹤鸣在 07 SpeechTTS
    assert "孙鹤鸣" in align_map["07"].matched_speakers

    # 张昊辰在 08 HiTranslationService
    assert "张昊辰" in align_map["08"].matched_speakers

    # 验证现场换序（Permutations）：现场 SpeechTTS (07) 发生于 SpeechASR (06) 之前
    chrono2 = res.chronological_alignments
    assert len(chrono2) == 9
    chrono_seqs = [c.item.seq for c in chrono2]
    assert chrono_seqs.index("07") < chrono_seqs.index("06")

def test_enhanced_state_machine_roadsign_and_buffer_isolation():
    """验证泛化性增强特性：
    1. 致谢语义消歧（答辩抗辩 vs. 真实收尾）；
    2. 主持人串场路标（精确人名/称谓/泛指）；
    3. 设备闲聊识别；
    4. 转场缓冲带隔离（闲聊与调麦不污染上一议题）。
    """
    from domains.meeting.tasks.agenda_minutes.alignment_engine import (
        _is_session_closing,
        _is_opening_signal,
        _is_equipment_or_chitchat,
        _detect_host_roadsign,
        align_agenda_with_transcript,
    )

    # 1. 致谢消歧
    assert _is_session_closing("整体结论是go，闭环之后再发，谢谢各位评委，好，拜拜") is True
    assert _is_session_closing("那我先下了，拜拜") is True
    assert _is_session_closing("谢谢评委提醒，我再解释一下原因") is False
    assert _is_session_closing("感谢老师提出的建议，我们后续版本会纳入考虑") is False

    # 2. 开场与设备闲聊
    assert _is_opening_signal("各位评委好，我来共享一下屏幕，开始汇报") is True
    assert _is_opening_signal("下面由我汇报本次版本") is True
    assert _is_equipment_or_chitchat("喂喂喂，能听到我声音吗？") is True
    assert _is_equipment_or_chitchat("大家稍等两分钟，我先倒杯水") is True
    assert _is_equipment_or_chitchat("这次版本主要优化了SpeechASR在离线长语音场景下的性能时延") is False

    # 3. 主持人交通警察路标
    plan = AgendaPlan(
        items=[
            AgendaItemParsed(seq="01", title="小艺慧记版本发布", presenters=["申家坤"]),
            AgendaItemParsed(seq="02", title="翻译海外需求评审", presenters=["刘畅"]),
            AgendaItemParsed(seq="03", title="SpeechASR评测", presenters=["陆敬怡", "林宇珂"]),
        ]
    )
    hubs = {"高雄", "徐锋"}

    # 主持人精准呼叫称谓（"林工"）
    seq, is_gen = _detect_host_roadsign("高雄", "辛苦刘工，下面有请林工汇报下一个议题", plan, hubs)
    assert seq == "03"
    assert is_gen is False

    # 主持人呼叫议题专名（"小艺慧记"）
    seq, is_gen = _detect_host_roadsign("高雄", "好的，我们切到下一个，小艺慧记", plan, hubs)
    assert seq == "01"
    assert is_gen is False

    # 主持人泛指串场
    seq, is_gen = _detect_host_roadsign("徐锋", "好的那有请下一位", plan, hubs)
    assert seq is None
    assert is_gen is True

    # 非主持人讨论内容不误报（如汇报人描述幻灯片布局"下面这个区域"）
    seq, is_gen = _detect_host_roadsign("刘畅", "然后小艺下面这个区域是配置项", plan, hubs)
    assert seq is None
    assert is_gen is False

    # 4. 转场缓冲带隔离测试：会间闲聊调设备不污染议题
    synthetic_transcript = """
刘畅 00:10:00
各位评委好，我来汇报翻译海外需求。本次核心是俄罗斯节点迁移与时延优化。
高雄 00:11:00
这个迁移会影响现网吗？
刘畅 00:11:20
不会，双机房热备，结论是go，闭环后发邮件，谢谢各位评委，好，拜拜。
高雄 00:12:00
好的辛苦刘畅。
李四 00:12:10
喂喂喂，能听到吗？大家稍微等两分钟，我去倒杯水。
高雄 00:14:00
下面有请林工汇报 SpeechASR。
林宇珂 00:14:20
各位评委晚上好，我来共享一下桌面，本次 SpeechASR 主要是 4.300 补丁版本。
高雄 00:15:00
时延指标怎么样？
林宇珂 00:15:20
时延下降 15%，整体结论是go，多谢大家拜拜。
"""
    res = align_agenda_with_transcript(plan, synthetic_transcript)
    assert res.discussed_count == 2
    a2 = [a for a in res.alignments if a.item.seq == "02"][0]
    a3 = [a for a in res.alignments if a.item.seq == "03"][0]

    # 验证议题 02 与 03 都被正常匹配且包含对应主讲人
    assert "刘畅" in a2.matched_speakers
    assert "林宇珂" in a3.matched_speakers

    # 核心验证：转场期间李四的倒水和调麦闲聊没有被贪婪污染到议题 02 的证据中！
    a2_evidence = a2.evidence_text
    assert "我去倒杯水" not in a2_evidence
    assert "喂喂喂" not in a2_evidence
    assert "俄罗斯节点" in a2_evidence

def test_alignment_ocr_table_mixed_discussed_and_skipped():
    """验证从 OCR 识别出的商评1议程单在对齐引擎中的表现。
    
    核心目标：
    - 议题 02（刘畅）：必须 discussed，且刘畅绝不能被误判为 Hub Speaker；
    - 议题 04（SpeechASR）：必须 discussed；
    - 议题 01 与 03：skipped。
    """
    ocr_shangping = """
**会议主题**：智慧域商用发布评审
**会议时间**：2026/06/15（周一）16:00-18:00（UTC+08:00）Beijing
- **全程与会人/主持人**：徐锋；索勋飞；吴友国；王昊；黄江；何瑄；左飞
- **分段列席人**：申家坤；方思邈；李腾飞；郑本令；沙彬斌；陈锴；刘畅；陈坤

| 序号 | 议题名称 | 汇报人/主讲人 | 预计时长 |
| :--- | :--- | :--- | :--- |
| 01 | 小艺慧记 CeliaMinutesService 1.4.5.500 版本发布商用版本 | 申家坤 | 15min |
| 02 | 翻译海外 HiTranslationService 21.1.1.300 商用版本发布 | 刘畅 | 15min |
| 03 | HAG 3.6.5.300 版本商用发布评审 | 沙彬斌；陈锴 | 15min |
| 04 | SpeechASR 1.4.5.302 商用版本评审 | 陆敬怡；林宇珂；赖朝辉 | 15min |
"""
    docx_path = AGENDA_DIR / "test1" / "商评1.docx"
    if not docx_path.is_file():
        pytest.skip("商评1.docx not found")

    doc = docx.Document(str(docx_path))
    transcript = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    plan = parse_agenda_text(ocr_shangping)

    res = align_agenda_with_transcript(plan, transcript)
    assert res.discussed_count == 2
    assert res.skipped_count == 2

    # 验证议题 02 成功对齐刘畅
    a2 = res.alignments[1]
    assert a2.item.seq == "02"
    assert a2.status == "discussed"
    assert "刘畅" in a2.matched_speakers
    assert len(a2.matched_blocks) >= 50

    # 验证议题 04 成功对齐陆敬怡
    a4 = res.alignments[3]
    assert a4.item.seq == "04"
    assert a4.status == "discussed"
    assert "陆敬怡" in a4.matched_speakers

def test_alignment_academic_forum_dual_presenters():
    """验证 test5 西工大（学术论坛模式）：
    - 组织主席张晓雷兼任议题04主讲人（双重身份仲裁）；
    - 转写中连续空行与多段发言顺利切块；
    - 4 项特邀报告全部有效对齐讨论。
    """
    docx_path = AGENDA_DIR / "test5" / "西工大会议.docx"
    if not docx_path.is_file():
        pytest.skip("西工大会议.docx not found")

    ocr_test5 = """
**会议名称**：CCF语音专委“走进高校”系列活动 第三站：西北工业大学
- **时间**：2022年11月05日 09:30-11:30
- **全程与会人/主持人**：张晓雷

| 序号 | 议题名称 | 汇报人/主讲人 | 预计时长 |
| :--- | :--- | :--- | :--- |
| 01 | 面向智能声学与临境通信应用的多通道声信号感知、处理与重构 | 陈景东 | 09:30-10:00 |
| 02 | 智能语音技术的新进展——NPU-ASLP Lab视角 | 谢磊 | 10:00-10:30 |
| 03 | 软硬联合优化的声学前端探索 | 付中华 | 10:30-11:00 |
| 04 | 鲁棒语音处理与安全 | 张晓雷 | 11:00-11:30 |
"""
    doc = docx.Document(str(docx_path))
    transcript = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    plan = parse_agenda_text(ocr_test5)

    res = align_agenda_with_transcript(plan, transcript)
    assert res.discussed_count == 4
    assert res.skipped_count == 0

    assert "陈景东" in res.alignments[0].matched_speakers
    assert "谢磊" in res.alignments[1].matched_speakers
    assert "付中华" in res.alignments[2].matched_speakers
    assert "张晓雷" in res.alignments[3].matched_speakers

def test_alignment_summit_keynotes_and_host_filtering():
    """验证 test6 学术年会（年会模式）：
    - 文本端自省发言人“主持人”并隔离串场；
    - 5 项特邀报告全部有效对齐讨论。
    """
    docx_path = AGENDA_DIR / "test6" / "学术年会.docx"
    if not docx_path.is_file():
        pytest.skip("学术年会.docx not found")

    ocr_test6 = """
- **全程与会人/主持人**：俞凯；高虹

| 序号 | 议题名称 | 汇报人/主讲人 | 预计时长 |
| :--- | :--- | :--- | :--- |
| 01 | 语音模仿与鉴别 | 陶建华 | 09:30-10:15 |
| 02 | 声学所语音识别最新进展及应用 | 颜永红 | 10:30-11:15 |
| 03 | 百度语音在建模技术、多模态和芯片方面的最新进展 | 贾磊 | 11:15-12:00 |
| 04 | 声纹识别应用的技术挑战 | 郑方 | 13:00-13:45 |
| 05 | 语音合成中的自然韵律多样性建模 | 俞凯 | 13:45-14:30 |
"""
    doc = docx.Document(str(docx_path))
    transcript = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    plan = parse_agenda_text(ocr_test6)

    res = align_agenda_with_transcript(plan, transcript)
    assert res.discussed_count == 5
    assert res.skipped_count == 0

    assert "陶建华" in res.alignments[0].matched_speakers
    assert "颜永红" in res.alignments[1].matched_speakers
    assert "贾磊" in res.alignments[2].matched_speakers
    assert "郑方" in res.alignments[3].matched_speakers
    assert "俞凯" in res.alignments[4].matched_speakers

def test_alignment_speaker_name_typo_fuzzy_reconciled():
    """验证 商评2 中 OCR 识别错字（陈啟错）的自动纠错与对齐：
    1. match_presenter_name 识别 错 与 锴 的 OCR 易混变体，打分达 0.95；
    2. parse_agenda_text 传入 transcript 时自动校准为真实发言人 陈啟锴；
    3. align_agenda_with_transcript 将第 5 项（HAG）成功判定为 discussed，不再沦为 [本次未讨论]。
    """
    from domains.meeting.tasks.agenda_minutes.agenda_parser import (
        match_presenter_name,
        parse_agenda_text,
    )
    from domains.meeting.tasks.agenda_minutes.alignment_engine import align_agenda_with_transcript

    # 1. 变体与形似容错打分
    assert match_presenter_name("陈啟错", "陈啟锴") >= 0.9
    assert match_presenter_name("陈启蒙", "陈啟锴") == 0.0  # 杜绝异人同姓字辈误伤

    docx_path = AGENDA_DIR / "test2" / "商评2.docx"
    txt_path = AGENDA_DIR / "test2" / "商评2.txt"
    if not docx_path.is_file() or not txt_path.is_file():
        pytest.skip("Test 2 files not found")

    doc = docx.Document(str(docx_path))
    transcript = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    # 模拟 OCR 出现错别字 陈啟错
    agenda_with_typo = txt_path.read_text(encoding="utf-8").replace("陈啟锴", "陈啟错")
    assert "陈啟错" in agenda_with_typo

    plan = parse_agenda_text(agenda_with_typo, transcript=transcript)
    hag_item = next(it for it in plan.items if "HAG" in it.title)
    assert "陈啟锴" in hag_item.presenters

    res = align_agenda_with_transcript(plan, transcript)
    hag_align = next(a for a in res.alignments if "HAG" in a.item.title)
    assert hag_align.status == "discussed"
    assert len(hag_align.matched_blocks) > 0
    assert res.discussed_count == 9
    assert res.skipped_count == 0

def test_alignment_anonymous_numbered_speaker_logs():
    """验证 test7 匿名说话人日志（Track B: Content & Roadsign Alignment）：
    1. group_transcript_blocks 成功解析 '发言者 1 00:00:03' 等全部 488 个发言块；
    2. 自适应路由嗅探识别出 Track B (CONTENT)；
    3. 状态机精准锚定 4 项已讨论议题与 1 项未讨论议题（01视频观看因未录制跳过）；
    4. 骨架后置强约束生成正确的时序重构与区间（00:00 ~ 00:18, 00:18 ~ 01:18 等）。
    """
    docx_path = AGENDA_DIR / "test7" / "说话人日志.docx"
    if not docx_path.is_file():
        pytest.skip("说话人日志.docx not found")

    import docx
    doc = docx.Document(docx_path)
    transcript = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    blocks = group_transcript_blocks(transcript)
    assert len(blocks) == 488
    assert blocks[0].speaker == "发言者 1"
    assert blocks[0].timestamp == "00:00:03"

    ocr_test7 = """
会议主题：何刚总北研所交流会
会议时间：2026年09月23日9:00—11:00

| 序号 | 时间 | 环节 | 内容说明 |
| :--- | :--- | :--- | :--- |
| 01 | 09:00-09:10 | 视频观看 | 活动开场视频播放 |
| 02 | 09:10-09:30 | 何刚总致辞 | 领导致辞 |
| 03 | 09:30-10:30 | 互动交流 | 现场问答、自由交流与讨论 |
| 04 | 10:30-10:40 | 任务令签署与授予 | 任务令签署仪式及授予环节 |
| 05 | 10:40-10:50 | 全体合影 | 全体参会人员合影留念 |
"""
    plan = parse_agenda_text(ocr_test7)
    res = align_agenda_with_transcript(plan, transcript)

    # 4 项讨论，1 项跳过
    assert res.discussed_count == 4
    assert res.skipped_count == 1

    align_map = {a.item.seq: a for a in res.alignments}

    # 01 视频观看 未播放，确定性置空
    assert align_map["01"].status == "skipped"
    assert len(align_map["01"].matched_blocks) == 0

    # 02 何刚总致辞 命中致辞路标
    assert align_map["02"].status == "discussed"
    assert len(align_map["02"].matched_blocks) > 30
    assert align_map["02"].start_index == 0

    # 03 互动交流 命中团队交流路标
    assert align_map["03"].status == "discussed"
    assert len(align_map["03"].matched_blocks) > 300

    # 04 任务令签署与授予 命中任务令签发仪式路标
    assert align_map["04"].status == "discussed"
    assert len(align_map["04"].matched_blocks) > 10

    # 05 全体合影 命中尾声合影留念路标
    assert align_map["05"].status == "discussed"
    assert len(align_map["05"].matched_blocks) >= 2

    # 验证时序重构与骨架锁定（纯会务动线与合影过场仪式已被程序化与智能过滤剔除）
    class DummyClient:
        pass

    agent = AgendaMinutesAgent(DummyClient())
    enforced = agent._enforce_agenda_invariants({}, res)
    items = enforced["agenda_items"]

    # 01 视频观看（0证据会务）与 05 全体合影（3块过场吆喝）已被干净剔除，保留3项实质研讨议题
    assert [it["agenda_seq"] for it in items] == ["02", "03", "04"]
    assert items[0]["agenda_title"] == "何刚总致辞"
    assert items[0]["time_range"] == "00:00 ~ 00:18"
    assert items[0]["discussion_state"] == "discussed"
    assert items[0]["status_tag"] == ""  # 非审批交流会，清空审批状态

    assert items[1]["agenda_title"] == "互动交流"
    assert items[1]["time_range"] == "00:18 ~ 01:18"
    assert items[1]["discussion_state"] == "discussed"
    assert items[1]["presenter"] == "何刚（答疑嘉宾）及现场参会团队"  # 方案 A：标准公文风
    assert items[1]["status_tag"] == ""

    assert items[2]["agenda_title"] == "任务令签署与授予"
    assert items[2]["time_range"] == "01:18 ~ 01:21"
    assert items[2]["discussion_state"] == "discussed"

    # 大盘统计动态同步更新为剔除过场后的实际议题数
    assert enforced["meeting_meta"]["agenda_stats"] == "既定议题共 3 项（有效审议 3 项 · 本次未讨论 0 项）"

    # 验证 Markdown 渲染：非审批交流会表格 100% 折叠为 3 列，且正文无「议题 02 · 」等生硬前缀
    md_out = format_agenda_minutes_markdown(enforced)
    assert "| 议题名称 | 汇报人 | 议题时长 |" in md_out
    assert "结论定调" not in md_out
    assert "| 何刚总致辞 | 何刚 | 00:00 ~ 00:18 |" in md_out
    assert "| 互动交流 | 何刚（答疑嘉宾）及现场参会团队 | 00:18 ~ 01:18 |" in md_out
    assert "### 何刚总致辞" in md_out
    assert "### 互动交流" in md_out
    assert "议题 02" not in md_out
    assert "议题 03" not in md_out

def test_ceremonial_and_non_agenda_filtering():
    """测试纯会务动线、作息日程、低密度合影过场的双轨过滤机制，同时确保严肃未讨论技术议题绝不误删。"""
    from domains.meeting.tasks.agenda_minutes.alignment_engine import (
        AgendaAlignment,
        AlignmentResult,
        DiscussionBlock,
        is_trivial_ceremonial_item,
    )
    from domains.meeting.tasks.agenda_minutes.agenda_parser import AgendaItemParsed, AgendaPlan

    # 1. 纯物理动线与生活作息：直接剔除
    item_bus = AgendaItemParsed(seq="01", title="乘车至大学城", presenters=[])
    align_bus = AgendaAlignment(item=item_bus, status="skipped")
    assert is_trivial_ceremonial_item(align_bus) is True

    item_walk = AgendaItemParsed(seq="02", title="漫步大学城", presenters=[])
    align_walk = AgendaAlignment(item=item_walk, status="skipped")
    assert is_trivial_ceremonial_item(align_walk) is True

    item_visit = AgendaItemParsed(seq="03", title="走进北科瑞声", presenters=[])
    align_visit = AgendaAlignment(item=item_visit, status="skipped")
    assert is_trivial_ceremonial_item(align_visit) is True

    item_lunch = AgendaItemParsed(seq="04", title="工作午餐", presenters=[])
    align_lunch = AgendaAlignment(item=item_lunch, status="discussed", matched_blocks=[
        DiscussionBlock(speaker="主持", timestamp="12:00", content="大家去餐厅就餐。")
    ])
    assert is_trivial_ceremonial_item(align_lunch) is True

    # 2. 零证据仪式（视频观看）：直接剔除，严禁显示为“本次未讨论”
    item_video = AgendaItemParsed(seq="05", title="开场短片观看", presenters=[])
    align_video = AgendaAlignment(item=item_video, status="skipped", matched_blocks=[])
    assert is_trivial_ceremonial_item(align_video) is True

    # 3. 低密度拍照过场（<=3块且<120字）：直接剔除
    item_photo = AgendaItemParsed(seq="06", title="全体合影", presenters=[])
    align_photo = AgendaAlignment(item=item_photo, status="discussed", matched_blocks=[
        DiscussionBlock(speaker="会务", timestamp="17:00", content="大家往前排站一下，看镜头。"),
        DiscussionBlock(speaker="主持", timestamp="17:01", content="好的三二一茄子。"),
    ])
    assert is_trivial_ceremonial_item(align_photo) is True

    # 4. 严肃技术/评审类议题即使未讨论（skipped），也坚决保留并标记为“本次未讨论”
    item_tech = AgendaItemParsed(seq="07", title="离线翻译模型轻量化算法评审", presenters=["刘畅"])
    align_tech = AgendaAlignment(item=item_tech, status="skipped", matched_blocks=[])
    assert is_trivial_ceremonial_item(align_tech) is False

    item_hag = AgendaItemParsed(seq="08", title="HAG 3.6.5.300版本商用发布评审", presenters=["沙彬斌"])
    align_hag = AgendaAlignment(item=item_hag, status="skipped", matched_blocks=[])
    assert is_trivial_ceremonial_item(align_hag) is False

    # 5. 验证管线在 _enforce_agenda_invariants 中完整剔除与跳过处理
    plan = AgendaPlan(items=[item_walk, item_photo, item_tech])
    res = AlignmentResult(
        plan=plan,
        alignments=[align_walk, align_photo, align_tech],
    )
    class DummyClient:
        pass
    agent = AgendaMinutesAgent(DummyClient())
    enforced = agent._enforce_agenda_invariants({}, res)
    enforced_items = enforced["agenda_items"]

    # 漫步大学城 与 全体合影 被过滤，仅保留 离线翻译模型轻量化算法评审
    assert len(enforced_items) == 1
    assert enforced_items[0]["agenda_seq"] == "07"
    assert enforced_items[0]["agenda_title"] == "离线翻译模型轻量化算法评审"
    assert enforced_items[0]["status_tag"] == "本次未讨论"
    assert enforced_items[0]["discussion_state"] == "skipped"
    assert enforced["meeting_meta"]["agenda_stats"] == "既定议题共 1 项（有效审议 0 项 · 本次未讨论 1 项）"

